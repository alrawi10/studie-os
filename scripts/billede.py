#!/usr/bin/env python3
"""Henter en figur fra pensum (PDF-side eller billede på PPTX-slide) til et Anki-kort.

Eksempler:
    # render side 12 af en PDF (kilden findes på navn i fagets manifest)
    billede.py --fag Cariologi --kilde "Chapter 2. Tooth tissues" --side 12
    # kun figuren (automatisk beskæring til billedet på siden)
    billede.py --fag Cariologi --kilde "Chapter 2" --side 6 --figur 1
    # beskær manuelt til en del af siden (brøkdele af bredde/højde: x0,y0,x1,y1)
    billede.py --fag Cariologi --kilde "Chapter 2" --side 12 --beskaer 0.1,0.4,0.9,0.8
    # største billede på slide 7 i en PPTX, og gem det direkte i Anki
    billede.py --fag Parodontologi --kilde "histopatologi" --side 7 --anki

Uden --anki gemmes PNG'en i .cache/billeder/ (så den kan ses igennem før brug).
Med --anki gemmes den også i Ankis mediemappe, og HTML til feltet `Billede` udskrives.

Brug:
    uv run --with pymupdf --with python-pptx --with pillow --with requests scripts/billede.py ...
"""
import argparse
import io
import json
import re
import sys
from pathlib import Path

ROD = Path(__file__).resolve().parent.parent
CACHE = ROD / ".cache" / "billeder"
ANKI_URL = "http://127.0.0.1:3141/"
MAKS_BREDDE = 1200


def normaliser(tekst):
    return re.sub(r"[^a-zæøå0-9]+", "", tekst.lower())


def find_kilde(fag, kilde):
    """Finder filen i fagets manifest ud fra (en del af) navnet; kilde kan også være en sti."""
    sti = Path(kilde)
    if sti.is_file():
        return sti
    kilder = ROD / "fag" / fag / "kilder"
    manifest = json.loads((kilder / ".manifest.json").read_text())
    soeg = normaliser(re.sub(r"\.(pdf|pptx|docx)$", "", kilde, flags=re.I))
    fundne = [p for p in manifest.values() if soeg in normaliser(p["navn"] or "")]
    fundne = [p for p in fundne if Path(p["sti"]).suffix.lower() in (".pdf", ".pptx")]
    if not fundne:
        sys.exit(f"Ingen PDF/PPTX i {fag} matcher '{kilde}'.")
    if len(fundne) > 1:
        # foretræk eksakt navn, ellers den korteste titel
        eksakt = [p for p in fundne if normaliser(re.sub(r"\.(pdf|pptx)$", "", p["navn"], flags=re.I)) == soeg]
        fundne = eksakt or sorted(fundne, key=lambda p: len(p["navn"]))
        andre = ", ".join(p["navn"] for p in fundne[1:4])
        if andre:
            print(f"(flere match – bruger '{fundne[0]['navn']}'; andre: {andre})", file=sys.stderr)
    return kilder / fundne[0]["sti"]


def fra_pdf(sti, side, beskaer, dpi, figur=None):
    import pymupdf as fitz

    with fitz.open(sti) as doc:
        if not 1 <= side <= doc.page_count:
            sys.exit(f"{sti.name} har {doc.page_count} sider – side {side} findes ikke.")
        page = doc[side - 1]
        r = page.rect
        clip = None
        if figur:
            # figurernes placering på siden (største først); luft over/under til figurtekst
            rects = [rc for img in page.get_images(full=True) for rc in page.get_image_rects(img[0])]
            rects = sorted({(rc.x0, rc.y0, rc.x1, rc.y1) for rc in rects if rc.width * rc.height > 0.01 * r.width * r.height},
                           key=lambda t: (t[2] - t[0]) * (t[3] - t[1]), reverse=True)
            if not rects:
                sys.exit(f"Side {side} i {sti.name} har ingen figurer – brug --beskaer eller hele siden.")
            if figur > len(rects):
                sys.exit(f"Side {side} har kun {len(rects)} figur(er).")
            x0, y0, x1, y1 = rects[figur - 1]
            luft_x, luft_y = 0.03 * r.width, 0.05 * r.height
            clip = fitz.Rect(max(r.x0, x0 - luft_x), max(r.y0, y0 - luft_y),
                             min(r.x1, x1 + luft_x), min(r.y1, y1 + luft_y))
        elif beskaer:
            x0, y0, x1, y1 = beskaer
            clip = fitz.Rect(r.x0 + x0 * r.width, r.y0 + y0 * r.height, r.x0 + x1 * r.width, r.y0 + y1 * r.height)
        return page.get_pixmap(dpi=dpi, clip=clip).tobytes("png")


def fra_pptx(sti, side, nr):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    slides = Presentation(sti).slides
    if not 1 <= side <= len(slides):
        sys.exit(f"{sti.name} har {len(slides)} slides – slide {side} findes ikke.")
    billeder = []

    def saml(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                saml(s.shapes)
            elif s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                billeder.append((s.width * s.height, s.image.blob))

    saml(slides[side - 1].shapes)
    if not billeder:
        sys.exit(f"Slide {side} i {sti.name} har ingen billeder (kun tekst/figurer tegnet i PowerPoint).")
    billeder.sort(key=lambda b: b[0], reverse=True)  # største først
    if nr > len(billeder):
        sys.exit(f"Slide {side} har kun {len(billeder)} billede(r).")
    return billeder[nr - 1][1]


def til_png(data):
    """Konverterer til PNG og begrænser bredden, så Anki-samlingen ikke bliver unødigt stor."""
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    if img.mode not in ("RGB", "RGBA"):
        img = img.convert("RGB")
    if img.width > MAKS_BREDDE:
        img = img.resize((MAKS_BREDDE, round(img.height * MAKS_BREDDE / img.width)), Image.LANCZOS)
    ud = io.BytesIO()
    img.save(ud, "PNG", optimize=True)
    return ud.getvalue()


def gem_i_anki(sti, filnavn):
    import requests

    r = requests.post(ANKI_URL, timeout=60,
                      headers={"Content-Type": "application/json", "Accept": "application/json, text/event-stream"},
                      json={"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                            "params": {"name": "store_media_file", "arguments": {"filename": filnavn, "path": str(sti)}}})
    tekst = r.content.decode("utf-8")
    svar = next((json.loads(l[5:]) for l in tekst.splitlines() if l.startswith("data:")), None) or r.json()
    if "error" in svar or svar.get("result", {}).get("isError"):
        sys.exit(f"Anki afviste billedet: {json.dumps(svar, ensure_ascii=False)[:300]}")
    return filnavn


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--fag", required=True)
    p.add_argument("--kilde", required=True, help="filnavn/titel (eller del af det) eller en sti")
    p.add_argument("--side", type=int, required=True, help="side (PDF) eller slide (PPTX), 1-indekseret")
    p.add_argument("--beskaer", help="x0,y0,x1,y1 som brøkdele af siden, fx 0,0.3,1,0.8 (kun PDF)")
    p.add_argument("--figur", type=int, help="PDF: beskær automatisk til figur nr. N på siden (1 = største)")
    p.add_argument("--nr", type=int, default=1, help="PPTX: hvilket billede på sliden (1 = største)")
    p.add_argument("--dpi", type=int, default=200)
    p.add_argument("--anki", action="store_true", help="gem i Ankis mediemappe og udskriv <img>-HTML")
    args = p.parse_args()

    kilde = find_kilde(args.fag, args.kilde)
    beskaer = tuple(float(v) for v in args.beskaer.split(",")) if args.beskaer else None
    if kilde.suffix.lower() == ".pdf":
        data = fra_pdf(kilde, args.side, beskaer, args.dpi, args.figur)
    elif kilde.suffix.lower() == ".pptx":
        data = fra_pptx(kilde, args.side, args.nr)
    else:
        sys.exit("Kun PDF og PPTX understøttes.")

    CACHE.mkdir(parents=True, exist_ok=True)
    slug = normaliser(kilde.stem)[:40] or "kilde"
    cfg_sti = ROD / "config" / "fag.json"
    studie = json.loads(cfg_sti.read_text()).get("studie", "studie") if cfg_sti.exists() else "studie"
    filnavn = f"{normaliser(studie)[:12] or 'studie'}_{args.fag.lower()}_{slug}_s{args.side}{'_c' if beskaer else ''}{f'_f{args.figur}' if args.figur else ''}{f'_b{args.nr}' if args.nr > 1 else ''}.png"
    filnavn = re.sub(r"[^a-z0-9_.-]", "", filnavn.replace("æ", "ae").replace("ø", "oe").replace("å", "aa"))
    ud = CACHE / filnavn
    ud.write_bytes(til_png(data))
    print(f"fil: {ud}")
    print(f"kilde: {kilde.name} · {'slide' if kilde.suffix.lower() == '.pptx' else 's.'} {args.side}")
    if args.anki:
        gem_i_anki(ud, filnavn)
        print(f'html: <img src="{filnavn}">')


if __name__ == "__main__":
    main()
