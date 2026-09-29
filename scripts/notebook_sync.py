#!/usr/bin/env python3
"""Uploader filer fra fag/<fag>/kilder/ til fagets NotebookLM-notebook via `nlm`.

- Opretter notebook'en, hvis config/fag.json ikke har et notebook_id endnu (og gemmer id'et).
- Springer filer over, der allerede ligger i notebook'en (sammenligning på normaliseret filnavn).
- Filer over 200 MB: PPTX/DOCX uploades som udtrukket tekst med slide-markører; store PDF'er springes over.
- Holder sig under MAKS_KILDER; ved for mange prioriteres forelæsninger/slides og pensumlitteratur.
- Skriver source-id tilbage i manifestet, så næste kørsel kun uploader nyt.

Brug:
    uv run --with python-pptx --with python-docx scripts/notebook_sync.py [--fag KOF] [--dry-run]
"""
import argparse
import json
import re
import subprocess
import sys
import time
from pathlib import Path

ROD = Path(__file__).resolve().parent.parent
NLM = str(Path.home() / ".local/bin/nlm")
MAKS_KILDER = 300  # NotebookLM Pro; sæt til 50 på gratisplanen
MAKS_BYTES = 200 * 1024 * 1024


def nlm(*args, forsoeg=3):
    for i in range(forsoeg):
        r = subprocess.run([NLM, *args], capture_output=True, text=True, timeout=900)
        if r.returncode == 0:
            return r.stdout
        besked = " ".join((r.stderr or r.stdout).split())
        if "timed out" not in besked.lower() or i == forsoeg - 1:
            raise RuntimeError(besked[:300] or "ukendt fejl")
        time.sleep(20 * (i + 1))


def nlm_json(*args):
    out = nlm(*args, "--json")
    start = min(i for i in (out.find("{"), out.find("[")) if i >= 0)
    return json.loads(out[start:])


def find_id(data):
    if isinstance(data, dict):
        for k in ("source_id", "notebook_id", "id"):
            if isinstance(data.get(k), str):
                return data[k]
        for v in data.values():
            if (fundet := find_id(v)):
                return fundet
    if isinstance(data, list):
        for v in data:
            if (fundet := find_id(v)):
                return fundet
    return None


def normaliser(titel):
    t = titel.lower().strip()
    t = re.sub(r"\s*\(tekst\)$", "", t)
    t = re.sub(r"\.(pdf|pptx|docx)$", "", t)
    t = re.sub(r"(\s*\(\d+\)|-\d)$", "", t)  # "fil (1)" / "fil-2" fra gentagne downloads
    return re.sub(r"[^a-zæøå0-9]+", "", t)


def prioritet(post):
    """Lavere = vigtigere. Forelæsninger/slides først, så pensumlitteratur, så resten."""
    tekst = f"{post['modul']} {post['navn']}".lower()
    if post["sti"].lower().endswith(".pptx") or re.search(r"forelæs|lecture|slide|\bf\d", tekst):
        return 0
    if re.search(r"pensum|litteratur|kapitel|chapter|kompendi|artikel", tekst):
        return 1
    return 2


def tekst_fra_fil(sti):
    if sti.suffix.lower() == ".pptx":
        from pptx import Presentation
        dele = []
        for i, slide in enumerate(Presentation(sti).slides, 1):
            tekst = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                tekst.append("Noter: " + slide.notes_slide.notes_text_frame.text)
            dele.append(f"--- Slide {i} ---\n" + "\n".join(tekst))
        return "\n\n".join(dele)
    if sti.suffix.lower() == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(sti).paragraphs)
    return None


def synk_fag(fag, info, cfg, dry_run):
    kilder = ROD / "fag" / fag / "kilder"
    manifest_sti = kilder / ".manifest.json"
    if not manifest_sti.exists():
        return [], [], []
    manifest = json.loads(manifest_sti.read_text())

    eksisterende = {}
    if not info.get("notebook_id"):
        if dry_run:
            print(f"  (ville oprette notebook '{info['notebook_titel']}')")
        else:
            info["notebook_id"] = find_id(nlm_json("notebook", "create", info["notebook_titel"]))
            # genindlæs før skrivning, så parallelle kørsler (--fag) ikke overskriver hinandens id'er
            cfg_sti = ROD / "config" / "fag.json"
            frisk = json.loads(cfg_sti.read_text())
            frisk["fag"][fag]["notebook_id"] = info["notebook_id"]
            cfg_sti.write_text(json.dumps(frisk, indent=2, ensure_ascii=False) + "\n")
            print(f"  oprettede notebook '{info['notebook_titel']}' ({info['notebook_id']})")
    if info.get("notebook_id"):
        kilder_nb = nlm_json("source", "list", info["notebook_id"])
        kilder_nb = kilder_nb.get("sources", kilder_nb) if isinstance(kilder_nb, dict) else kilder_nb
        eksisterende = {normaliser(s["title"]): s["id"] for s in kilder_nb}

    ventende = [(k, p) for k, p in manifest.items() if not p.get("notebooklm_source_id")]
    ventende.sort(key=lambda kp: prioritet(kp[1]))
    plads = MAKS_KILDER - len(eksisterende)

    uploadet, fundet, udeladt = [], [], []
    for noegle, post in ventende:
        if (sid := eksisterende.get(normaliser(post["navn"]))) and not post.get("notebooklm_forrige_source_id"):
            post["notebooklm_source_id"] = sid
            fundet.append(post["sti"])
            continue
        sti = kilder / post["sti"]
        if plads <= 0:
            udeladt.append(f"{post['sti']} (notebook fuld)")
            continue
        stor = sti.stat().st_size > MAKS_BYTES
        if dry_run:
            uploadet.append(post["sti"] + (" [som tekst]" if stor else ""))
            plads -= 1
            continue
        try:
            if stor:
                tekst = tekst_fra_fil(sti)
                if not tekst:
                    udeladt.append(f"{post['sti']} (over 200 MB, kan ikke konverteres)")
                    continue
                res = nlm_json("source", "add", info["notebook_id"], "--text", tekst, "--title", f"{post['navn']} (tekst)")
            else:
                res = nlm_json("source", "add", info["notebook_id"], "--file", str(sti), "--title", post["navn"])
        except (RuntimeError, subprocess.TimeoutExpired, json.JSONDecodeError, ValueError) as e:
            udeladt.append(f"{post['sti']} (fejl: {str(e)[:120]})")
            continue
        post["notebooklm_source_id"] = find_id(res) or "ukendt"
        uploadet.append(post["sti"] + (" [som tekst]" if stor else ""))
        plads -= 1
        manifest_sti.write_text(json.dumps(manifest, indent=2, ensure_ascii=False))  # gem løbende

    if not dry_run:
        manifest_sti.write_text(json.dumps(manifest, indent=2, ensure_ascii=False))
    return uploadet, fundet, udeladt


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--fag")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()

    cfg = json.loads((ROD / "config" / "fag.json").read_text())
    valgte = [args.fag] if args.fag else list(cfg["fag"])
    for fag in valgte:
        print(f"\n== {fag}")
        uploadet, fundet, udeladt = synk_fag(fag, cfg["fag"][fag], cfg, args.dry_run)
        print(f"  uploadet: {len(uploadet)} · fandtes allerede: {len(fundet)} · udeladt: {len(udeladt)}")
        for linje in uploadet:
            print("  + " + linje)
        for linje in udeladt:
            print("  - " + linje)


if __name__ == "__main__":
    sys.exit(main())
